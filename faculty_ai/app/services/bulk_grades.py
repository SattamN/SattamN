"""Bulk grade updates via a re-uploaded Excel file (the counterpart to build_editable_xlsx).

Flow: instructor downloads the editable gradebook -> edits mark cells in Excel -> re-uploads.
We always compute a diff first (commit=False); nothing is written until the instructor confirms
with commit=True, mirroring the same review-before-write discipline used everywhere else in this
system (rubric approval, review-request decisions, etc.).
"""
import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Optional

from openpyxl import load_workbook
from sqlalchemy.orm import Session, selectinload

from ..models import Exam, Grade, Question, Submission, User
from .audit import audit

TOL = 1e-6
ID_ALIASES = {"university id", "university_id", "student id", "id", "الرقم الجامعي", "رقم الطالب", "الرقم"}
Q_RE = re.compile(r"^\s*Q\s*(\d+)", re.IGNORECASE)


@dataclass
class ParsedRow:
    university_id: str
    marks: dict[int, float]  # question number -> new mark
    row_number: int


@dataclass
class ParseResult:
    rows: list[ParsedRow]
    warnings: list[str] = field(default_factory=list)


def _clean_id(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def parse_editable_xlsx(data: bytes) -> ParseResult:
    """Reads the workbook produced by build_editable_xlsx (or any sheet with a University ID column
    and Q<number> columns). Tolerant of the instructions row being present or removed."""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"Could not read this file as an Excel workbook: {e}") from e
    ws = wb.active
    raw = [list(r) for r in ws.iter_rows(values_only=True)]
    raw = [r for r in raw if any(c not in (None, "") for c in r)]
    if not raw:
        raise ValueError("The file is empty")

    header_idx = None
    for i, row in enumerate(raw[:5]):  # header is row 1 or 2 (instructions row may or may not be present)
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        if any(c in ID_ALIASES for c in cells):
            header_idx = i
            break
    if header_idx is None:
        raise ValueError("Could not find a 'University ID' column in the first few rows")

    header = raw[header_idx]
    id_col = next(i for i, c in enumerate(header) if c and str(c).strip().lower() in ID_ALIASES)
    q_cols: dict[int, int] = {}
    for i, c in enumerate(header):
        if not c:
            continue
        m = Q_RE.match(str(c))
        if m:
            q_cols[int(m.group(1))] = i
    if not q_cols:
        raise ValueError("Could not find any question columns (expected headers like 'Q1', 'Q2 (max 10)', ...)")

    rows, warnings = [], []
    for r_offset, r in enumerate(raw[header_idx + 1 :], start=header_idx + 2):
        uid = _clean_id(r[id_col] if id_col < len(r) else None)
        if not uid:
            continue  # blank separator row, ignore silently
        marks: dict[int, float] = {}
        for qnum, ci in q_cols.items():
            val = r[ci] if ci < len(r) else None
            if val is None or val == "":
                continue
            try:
                marks[qnum] = float(val)
            except (TypeError, ValueError):
                warnings.append(f"Row {r_offset}, Q{qnum}: '{val}' is not a number, skipped")
        if marks:
            rows.append(ParsedRow(university_id=uid, marks=marks, row_number=r_offset))
    if not rows:
        raise ValueError("No usable rows found (every row was blank or unreadable)")
    return ParseResult(rows=rows, warnings=warnings)


# --------------------------------------------------------------------------- diff


def compute_diff(db: Session, exam: Exam, parsed: ParseResult) -> dict:
    questions = {q.number: q for q in exam.questions}
    subs = (
        db.query(Submission)
        .options(selectinload(Submission.grades), selectinload(Submission.student))
        .filter(Submission.exam_id == exam.id)
        .all()
    )
    by_uid = {s.student.university_id: s for s in subs}

    changes, warnings, unmatched = [], list(parsed.warnings), []
    seen_uids = set()
    for row in parsed.rows:
        if row.university_id in seen_uids:
            warnings.append(f"Row {row.row_number}: duplicate University ID {row.university_id}, later row wins")
        seen_uids.add(row.university_id)

        sub = by_uid.get(row.university_id)
        if sub is None:
            unmatched.append({"row": row.row_number, "university_id": row.university_id,
                              "reason": "No submission for this student in this exam"})
            continue
        by_qid = {g.question_id: g for g in sub.grades}
        for qnum, new_mark in row.marks.items():
            q = questions.get(qnum)
            if q is None:
                warnings.append(f"Row {row.row_number}: Q{qnum} does not exist in this exam, skipped")
                continue
            if new_mark < -TOL or new_mark > q.max_mark + TOL:
                warnings.append(f"Row {row.row_number}: Q{qnum} = {new_mark:g} is outside 0–{q.max_mark:g}, skipped")
                continue
            grade = by_qid.get(q.id)
            if grade is None:
                warnings.append(f"Row {row.row_number}: Q{qnum} has not been graded yet for this student, skipped")
                continue
            current = grade.final_mark if grade.final_mark is not None else grade.ai_mark
            if current is not None and abs(current - new_mark) < TOL:
                continue  # no real change
            changes.append({
                "submission_id": sub.id, "university_id": row.university_id, "student_name": sub.student.name,
                "question_id": q.id, "question_number": qnum, "max_mark": q.max_mark,
                "current_mark": current, "new_mark": round(new_mark, 4),
                "currently_approved": sub.status == "approved",
            })
    return {"changes": changes, "warnings": warnings, "unmatched": unmatched}


# --------------------------------------------------------------------------- apply


def apply_diff(db: Session, exam: Exam, diff: dict, user: User) -> dict:
    """Writes every change in `diff['changes']` (as returned by compute_diff). Reopens an approved
    submission before editing it, exactly like the existing manual-edit rule requires, and audits
    every write -- same shape as a single PATCH /grades/{id}."""
    now = datetime.now(timezone.utc)
    reopened, applied = set(), 0
    grade_ids = [c["question_id"] for c in diff["changes"]]
    subs_by_id = {
        s.id: s
        for s in db.query(Submission).options(selectinload(Submission.grades)).filter(
            Submission.id.in_({c["submission_id"] for c in diff["changes"]})
        ).all()
    } if diff["changes"] else {}

    for c in diff["changes"]:
        sub = subs_by_id[c["submission_id"]]
        if sub.status == "approved" and sub.id not in reopened:
            sub.status, sub.approved_at = "graded", None
            audit(db, user, "submission_reopened", "submission", sub.id, {"reason": "bulk_excel_import"})
            reopened.add(sub.id)
        grade = next(g for g in sub.grades if g.question_id == c["question_id"])
        before = grade.final_mark
        grade.final_mark = c["new_mark"]
        grade.decision = "accepted_ai" if (grade.ai_mark is not None and abs(grade.ai_mark - c["new_mark"]) < TOL) else "edited"
        grade.reviewer_note = "Bulk Excel import" if not grade.reviewer_note else f"{grade.reviewer_note} | Bulk Excel import"
        audit(db, user, "grade_set", "grade", grade.id, {
            "submission_id": sub.id, "question": c["question_number"], "ai_mark": grade.ai_mark,
            "before": before, "after": grade.final_mark, "decision": grade.decision, "source": "bulk_excel_import",
        })
        applied += 1

    audit(db, user, "bulk_excel_import_committed", "exam", exam.id,
          {"applied": applied, "reopened": len(reopened), "warnings": len(diff["warnings"]), "unmatched": len(diff["unmatched"])})
    db.commit()
    return {"applied": applied, "reopened_submissions": len(reopened)}
