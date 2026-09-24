"""Parse a class roster (.xlsx / .csv) with English or Arabic headers."""
import csv
import io
from pathlib import Path

ALIASES = {
    "university_id": {"university id", "university_id", "student id", "id", "الرقم الجامعي", "رقم الطالب", "الرقم"},
    "name": {"student name", "name", "full name", "اسم الطالب", "الاسم"},
    "email": {"email", "e-mail", "البريد", "البريد الإلكتروني", "البريد الالكتروني"},
    "section": {"section", "الشعبة", "شعبة"},
}
MAX_ROWS = 2000


def _canon(header: str) -> str | None:
    h = (header or "").strip().lower()
    for key, names in ALIASES.items():
        if h in names:
            return key
    return None


def _clean_id(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # Excel turns 441234567 into 441234567.0
    return str(v).strip()


def parse_roster(filename: str, data: bytes) -> tuple[list[dict], list[str]]:
    """Returns (rows, warnings). Each row: university_id, name, email, section."""
    ext = Path(filename or "").suffix.lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
    elif ext == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        raw = list(csv.reader(io.StringIO(text)))
    else:
        raise ValueError("Roster must be an .xlsx or .csv file")

    raw = [r for r in raw if any(c not in (None, "") for c in r)]
    if not raw:
        raise ValueError("The roster file is empty")
    cols = {i: _canon(str(h)) for i, h in enumerate(raw[0])}
    if "university_id" not in cols.values():
        raise ValueError("Missing a 'University ID' column (or 'الرقم الجامعي')")
    if len(raw) - 1 > MAX_ROWS:
        raise ValueError(f"Too many rows (max {MAX_ROWS})")

    rows, warnings, seen = [], [], set()
    for n, r in enumerate(raw[1:], start=2):
        rec = {"university_id": "", "name": "", "email": "", "section": ""}
        for i, key in cols.items():
            if key and i < len(r):
                rec[key] = _clean_id(r[i]) if key in ("university_id", "section") else str(r[i] or "").strip()
        if not rec["university_id"]:
            warnings.append(f"Row {n}: no university ID, skipped")
        elif rec["university_id"] in seen:
            warnings.append(f"Row {n}: duplicate ID {rec['university_id']}, skipped")
        else:
            seen.add(rec["university_id"])
            rows.append(rec)
    return rows, warnings
