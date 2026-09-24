"""Tests for the grading core (no network: the Anthropic client is faked)."""
import json
import io
import os
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

os.environ["DATA_DIR"] = tempfile.mkdtemp(prefix="faculty_ai_test_")
os.environ.pop("ANTHROPIC_API_KEY", None)
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from sqlalchemy import create_engine

from app import models
from app.config import get_settings
from app.database import Base, configure_sqlite, get_db
from app.main import app
from app.models import AICallLog, AuditLog, Course, Exam, Grade, Question, Student, Submission, User
from app.routers import auth as auth_router
from app.security import hash_password
from app.schemas import QuestionResult
from app.services import ai_grader as ag
from app.services.ai_grader import AIGrader, evaluate_review_flags, get_grader, grade_submission

PASSWORD = "Passw0rd!x"
RUBRIC = {"criteria": [
    {"criterion": "Correct method", "max_mark": 3, "description": ""},
    {"criterion": "Calculations", "max_mark": 4, "description": ""},
    {"criterion": "Final answer", "max_mark": 3, "description": ""},
]}


# ------------------------------------------------------------------ fakes / helpers


class FakeMessages:
    def __init__(self, outputs):
        self.outputs = list(outputs)
        self.calls = []

    def create(self, **kw):
        self.calls.append(kw)
        out = self.outputs.pop(0)
        block = SimpleNamespace(type="tool_use", name=kw["tools"][0]["name"], input=out)
        return SimpleNamespace(content=[block], stop_reason="tool_use",
                               usage=SimpleNamespace(input_tokens=1200, output_tokens=300))


class FakeClient:
    def __init__(self, outputs):
        self.messages = FakeMessages(outputs)


def item(n, mark=7.5, mx=10, conf=0.91, needs_review=False, status="answered", breakdown="ok", reason="Correct method, arithmetic slip."):
    if breakdown == "ok":
        breakdown = [
            {"criterion": "Correct method", "mark": 3, "max_mark": 3},
            {"criterion": "Calculations", "mark": mark - 5 if mark >= 5 else 0, "max_mark": 4},
            {"criterion": "Final answer", "mark": 2 if mark >= 5 else 0, "max_mark": 3},
        ]
    return {"question_number": n, "student_answer_transcription": "M = wL^2/8 = 45 kNm", "answer_status": status,
            "mark": mark, "max_mark": mx, "reason": reason, "confidence": conf, "needs_review": needs_review,
            "rubric_breakdown": breakdown}


def make_question(number=1, max_mark=10, rubric=RUBRIC):
    q = Question(id=number, exam_id=1, number=number, max_mark=max_mark, text="Q", rubric_json=rubric, rubric_approved=True)
    return q


def qr(**kw) -> QuestionResult:
    return QuestionResult.model_validate(item(1, **kw))


@pytest.fixture()
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    configure_sqlite(engine)
    Base.metadata.create_all(engine)
    S = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    session = S()
    session.factory = S
    yield session
    session.close()


def seed(db, n_questions=2):
    fac = User(university_id="F1", name="Dr. Ahmed", role="faculty", password_hash="x", must_change_password=False)
    db.add(fac)
    db.flush()
    course = Course(name="Reinforced Concrete", code="CE 301", clos=[], instructor_id=fac.id)
    exam = Exam(course=course, title="Midterm")
    for i in range(1, n_questions + 1):
        exam.questions.append(Question(number=i, text=f"Question {i}", max_mark=10, model_answer="ans",
                                       rubric_json=RUBRIC, rubric_approved=True))
    student = Student(name="أحمد الشمري", university_id="441001")
    db.add_all([course, exam, student])
    db.flush()
    f = Path(get_settings().uploads_dir) / "t.png"
    f.parent.mkdir(parents=True, exist_ok=True)
    f.write_bytes(b"\x89PNG\r\n\x1a\nfake")
    sub = Submission(exam_id=exam.id, student_id=student.id, file_paths=["uploads/t.png"], status="uploaded")
    db.add(sub)
    db.commit()
    return exam, sub


# ------------------------------------------------------------------ review rules


def test_clean_result_has_no_flags():
    assert evaluate_review_flags(qr(), make_question(), 0.75) == []


def test_confidence_is_only_one_signal():
    q = make_question()
    assert ag.FLAG_LOW_CONFIDENCE in evaluate_review_flags(qr(conf=0.6), q, 0.75)
    # high confidence but inconsistent breakdown must still be flagged
    bad = qr(conf=0.99, breakdown=[{"criterion": "Correct method", "mark": 3, "max_mark": 3},
                                   {"criterion": "Calculations", "mark": 4, "max_mark": 4},
                                   {"criterion": "Final answer", "mark": 3, "max_mark": 3}])  # sums to 10, mark says 7.5
    assert ag.FLAG_BREAKDOWN_SUM in evaluate_review_flags(bad, q, 0.75)


def test_other_rules():
    q = make_question()
    assert ag.FLAG_MODEL_FLAGGED in evaluate_review_flags(qr(needs_review=True), q, 0.75)
    assert ag.FLAG_ANSWER_UNCLEAR in evaluate_review_flags(qr(status="unclear"), q, 0.75)
    assert ag.FLAG_BLANK_WITH_MARKS in evaluate_review_flags(qr(status="blank"), q, 0.75)
    assert ag.FLAG_MARK_OUT_OF_RANGE in evaluate_review_flags(qr(mark=12), q, 0.75)
    assert ag.FLAG_MAX_MISMATCH in evaluate_review_flags(qr(mx=8), q, 0.75)
    assert ag.FLAG_EMPTY_REASON in evaluate_review_flags(qr(reason=" "), q, 0.75)
    # criteria not matching the approved rubric / missing breakdown
    assert ag.FLAG_RUBRIC_MISMATCH in evaluate_review_flags(qr(breakdown=[]), q, 0.75)
    renamed = qr(breakdown=[{"criterion": "Something else", "mark": 7.5, "max_mark": 10}])
    assert ag.FLAG_RUBRIC_MISMATCH in evaluate_review_flags(renamed, q, 0.75)


# ------------------------------------------------------------------ grading pipeline


def test_grade_submission_persists_ai_and_keeps_final_separate(db):
    exam, sub = seed(db)
    client = FakeClient([{"questions": [item(1), item(2, conf=0.5, mark=12)]}])  # Q2: low conf + out of range
    grade_submission(db, sub, grader=AIGrader(client=client))

    db.refresh(sub)
    assert sub.status == "graded"
    g1, g2 = sorted(sub.grades, key=lambda g: g.question_id)
    assert g1.ai_mark == 7.5 and g1.final_mark is None and g1.decision == "pending"
    assert g1.needs_review is False and g1.review_flags == []
    assert g2.ai_mark == 10  # clamped for storage ...
    assert json.loads(g2.ai_raw_response)["mark"] == 12  # ... but the raw output is preserved
    assert set(g2.review_flags) >= {ag.FLAG_LOW_CONFIDENCE, ag.FLAG_MARK_OUT_OF_RANGE} and g2.needs_review
    assert g1.ai_answer_text.startswith("M =")

    # audit log + forced structured output
    logs = db.query(AICallLog).all()
    assert len(logs) == 1 and logs[0].input_tokens == 1200 and json.loads(logs[0].raw_response)["questions"]
    call = client.messages.calls[0]
    assert call["tool_choice"] == {"type": "tool", "name": "submit_grades"}
    assert "UNTRUSTED" in call["system"]  # prompt-injection guard present
    assert g1.created_at is not None and g1.updated_at is not None


def test_missing_question_is_flagged(db):
    exam, sub = seed(db)
    grade_submission(db, sub, grader=AIGrader(client=FakeClient([{"questions": [item(1)]}])))
    g = {g.question_id: g for g in sub.grades}
    q2 = next(q for q in exam.questions if q.number == 2)
    assert g[q2.id].ai_mark is None and ag.FLAG_MISSING in g[q2.id].review_flags


def test_invalid_output_is_retried_and_both_calls_logged(db):
    exam, sub = seed(db, 1)
    client = FakeClient([{"questions": [{"question_number": 1}]}, {"questions": [item(1)]}])
    grade_submission(db, sub, grader=AIGrader(client=client))
    assert sub.status == "graded" and len(client.messages.calls) == 2
    logs = db.query(AICallLog).order_by(AICallLog.id).all()
    assert logs[0].error.startswith("validation_failed") and logs[1].error is None
    assert "rejected" in json.dumps(client.messages.calls[1]["messages"])


def test_grading_requires_approved_rubrics_and_protects_reviewed_work(db):
    exam, sub = seed(db, 1)
    exam.questions[0].rubric_approved = False
    db.commit()
    with pytest.raises(ValueError):
        grade_submission(db, sub, grader=AIGrader(client=FakeClient([])))
    exam.questions[0].rubric_approved = True
    db.commit()
    grade_submission(db, sub, grader=AIGrader(client=FakeClient([{"questions": [item(1)]}])))
    sub.grades[0].decision, sub.grades[0].final_mark = "edited", 6
    db.commit()
    with pytest.raises(ag.GradingConflict):
        grade_submission(db, sub, grader=AIGrader(client=FakeClient([])))


def test_ai_failure_marks_submission_failed(db):
    exam, sub = seed(db, 1)
    bad = {"questions": [{"question_number": 1}]}
    with pytest.raises(ag.AIError):
        grade_submission(db, sub, grader=AIGrader(client=FakeClient([bad, bad])))
    db.refresh(sub)
    assert sub.status == "failed" and "validation" in sub.error


def test_run_exam_grading_concurrent(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path/'t.db'}", connect_args={"check_same_thread": False, "timeout": 30})
    configure_sqlite(engine)
    Base.metadata.create_all(engine)
    S = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    db = S()
    exam, sub = seed(db, 1)
    subs = [sub]
    for i in range(4):
        st = Student(name=f"S{i}", university_id=f"9{i}")
        db.add(st); db.flush()
        s = Submission(exam_id=exam.id, student_id=st.id, file_paths=["uploads/t.png"], status="queued")
        db.add(s); subs.append(s)
    db.commit()
    ids = [s.id for s in subs]
    client = FakeClient([{"questions": [item(1)]} for _ in ids])
    result = ag.run_exam_grading(ids, grader=AIGrader(client=client), session_factory=S)
    assert all(v is None for v in result.values()), result
    assert {s.status for s in S().query(Submission).all()} == {"graded"}


# ------------------------------------------------------------------ API: review -> approve -> gradebook -> export


@pytest.fixture()
def client_api():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    configure_sqlite(engine)
    Base.metadata.create_all(engine)
    S = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)

    def _db():
        d = S()
        try:
            yield d
        finally:
            d.close()

    seed_db = S()
    pw = hash_password(PASSWORD)
    for uid, role in (("F1", "faculty"), ("F2", "faculty"), ("A1", "admin")):
        seed_db.add(User(university_id=uid, name=f"User {uid}", role=role, password_hash=pw, must_change_password=False))
    seed_db.commit()
    seed_db.close()

    fake = FakeClient([])
    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[get_grader] = lambda: AIGrader(client=fake)
    auth_router._failures.clear()
    api = TestClient(app)

    def login(uid, password=PASSWORD):
        r = api.post("/auth/login", json={"university_id": uid, "password": password})
        assert r.status_code == 200, r.text
        return {"Authorization": f"Bearer {r.json()['access_token']}"}

    api.login = login
    api.S = S
    api.headers.update(login("F1"))  # default: faculty F1
    yield api, fake
    app.dependency_overrides.clear()


def test_full_api_flow(client_api):
    api, fake = client_api
    cid = api.post("/courses", json={"name": "Reinforced Concrete", "code": "CE 301"}).json()["id"]
    eid = api.post(f"/courses/{cid}/exams", data={"title": "Midterm"}).json()["id"]
    q_ids = []
    for n in (1, 2):
        r = api.post(f"/exams/{eid}/questions", json={"number": n, "text": f"Q{n}", "max_mark": 10, "model_answer": "x"})
        assert r.status_code == 201
        q_ids.append(r.json()["id"])

    # grading is blocked until rubrics are approved
    files = {"files": ("a.png", b"\x89PNGfake", "image/png")}
    sid = api.post(f"/exams/{eid}/submissions", data={"university_id": "441001", "student_name": "أحمد"}, files=files).json()["id"]
    assert api.post(f"/submissions/{sid}/grade").status_code == 409

    # rubric must sum to the question mark
    bad = {"criteria": [{"criterion": "A", "max_mark": 3}]}
    assert api.put(f"/questions/{q_ids[0]}/rubric", json=bad).status_code == 400
    for qid in q_ids:
        api.post(f"/questions/{qid}/rubric/generate?mode=default")
    assert api.post(f"/exams/{eid}/rubrics/approve-all").json()["approved"] == 2

    crit = api.get(f"/exams/{eid}/questions").json()[0]["rubric"]["criteria"]
    good = lambda mark_scale=1: [{"criterion": c["criterion"], "mark": c["max_mark"] * mark_scale, "max_mark": c["max_mark"]} for c in crit]
    q1 = item(1, mark=10, conf=0.95, breakdown=good())
    q2 = item(2, mark=5, conf=0.4, breakdown=good(0.5))  # low confidence -> must be reviewed
    fake.messages.outputs.append({"questions": [q1, q2]})
    out = api.post(f"/submissions/{sid}/grade").json()
    assert out["status"] == "graded" and out["n_pending_review"] == 1

    # cannot approve while a flagged question is unreviewed
    assert api.post(f"/submissions/{sid}/approve").status_code == 409
    review = api.get(f"/submissions/{sid}/review").json()
    flagged = next(g for g in review["grades"] if g["needs_review"])
    assert api.patch(f"/grades/{flagged['id']}", json={"final_mark": 11}).status_code == 400
    r = api.patch(f"/grades/{flagged['id']}", json={"final_mark": 6.5, "reviewer_note": "gave credit for alt. method"})
    assert r.json()["decision"] == "edited" and r.json()["ai_mark"] == 5  # AI mark untouched

    assert api.post(f"/submissions/{sid}/approve").json()["status"] == "approved"
    assert api.patch(f"/grades/{flagged['id']}", json={"final_mark": 7}).status_code == 409  # locked until reopened

    gb = api.get(f"/exams/{eid}/gradebook").json()
    assert gb["rows"][0]["total"] == 16.5 and gb["total_max"] == 20 and gb["stats"]["count"] == 1
    assert gb["rows"][0]["marks"] == {"1": 10.0, "2": 6.5}

    # exports (Arabic student name must survive all formats)
    x = api.get(f"/exams/{eid}/export/xlsx")
    assert x.status_code == 200
    import io, openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(x.content))["Gradebook"]
    assert ws["B3"].value == "أحمد" and ws["C3"].value == 10 and str(ws["E3"].value).startswith("=SUM")
    c = api.get(f"/exams/{eid}/export/csv")
    assert c.status_code == 200 and "أحمد" in c.content.decode("utf-8-sig")
    p = api.get(f"/exams/{eid}/export/pdf")
    assert p.status_code == 200 and p.content.startswith(b"%PDF")


def test_bulk_approve_only_unflagged(client_api):
    api, fake = client_api
    cid = api.post("/courses", json={"name": "C"}).json()["id"]
    eid = api.post(f"/courses/{cid}/exams", data={"title": "E"}).json()["id"]
    qid = api.post(f"/exams/{eid}/questions", json={"number": 1, "text": "Q", "max_mark": 10}).json()["id"]
    api.post(f"/questions/{qid}/rubric/generate?mode=default")
    api.post(f"/exams/{eid}/rubrics/approve-all")
    crit = api.get(f"/exams/{eid}/questions").json()[0]["rubric"]["criteria"]
    bd = lambda: [{"criterion": c["criterion"], "mark": c["max_mark"], "max_mark": c["max_mark"]} for c in crit]
    ids = []
    for uid, conf in (("1", 0.95), ("2", 0.3)):
        sid = api.post(f"/exams/{eid}/submissions", data={"university_id": uid},
                       files={"files": (f"{uid}.png", b"\x89PNGx", "image/png")}).json()["id"]
        fake.messages.outputs.append({"questions": [item(1, mark=10, conf=conf, breakdown=bd())]})
        api.post(f"/submissions/{sid}/grade")
        ids.append(sid)
    assert api.post(f"/exams/{eid}/approve-unflagged").json() == {"approved": 1, "held_for_review": 1}
    statuses = {s["id"]: s["status"] for s in api.get(f"/exams/{eid}/submissions").json()}
    assert statuses == {ids[0]: "approved", ids[1]: "graded"}


def test_upload_validation(client_api):
    api, _ = client_api
    cid = api.post("/courses", json={"name": "C"}).json()["id"]
    eid = api.post(f"/courses/{cid}/exams", data={"title": "E"}).json()["id"]
    r = api.post(f"/exams/{eid}/submissions", data={"university_id": "1"}, files={"files": ("x.exe", b"MZ", "application/octet-stream")})
    assert r.status_code == 400
    r = api.post(f"/exams/{eid}/grade")
    assert r.status_code == 409  # no questions yet


# ------------------------------------------------------------------ accounts, roles, student portal


def _graded_exam(api, fake, student_ids, ai_mark=8.0, conf=0.95):
    """Faculty F1 (default headers): course + exam with one approved-rubric question, papers uploaded & AI-graded."""
    cid = api.post("/courses", json={"name": "Reinforced Concrete", "code": "CE 302"}).json()["id"]
    eid = api.post(f"/courses/{cid}/exams", data={"title": "Midterm"}).json()["id"]
    qid = api.post(f"/exams/{eid}/questions", json={"number": 1, "text": "Q", "max_mark": 10}).json()["id"]
    api.post(f"/questions/{qid}/rubric/generate?mode=default")
    api.post(f"/exams/{eid}/rubrics/approve-all")
    crit = api.get(f"/exams/{eid}/questions").json()[0]["rubric"]["criteria"]
    parts, left = [], ai_mark
    for c in crit:
        m = min(c["max_mark"], left)
        parts.append({"criterion": c["criterion"], "mark": m, "max_mark": c["max_mark"]})
        left -= m
    sids = {}
    for uid in student_ids:
        sid = api.post(f"/exams/{eid}/submissions", data={"university_id": uid},
                       files={"files": (f"{uid}.png", b"\x89PNGx", "image/png")}).json()["id"]
        fake.messages.outputs.append({"questions": [item(1, mark=ai_mark, conf=conf, breakdown=parts)]})
        assert api.post(f"/submissions/{sid}/grade").status_code == 200
        sids[uid] = sid
    return cid, eid, sids


def test_requires_login_and_role(client_api):
    api, _ = client_api
    anon = TestClient(app)
    assert anon.get("/courses").status_code == 401
    assert anon.post("/auth/login", json={"university_id": "F1", "password": "wrong"}).status_code == 401
    # a student token cannot use faculty endpoints
    cid = api.post("/courses", json={"name": "C"}).json()["id"]
    csv_bytes = "University ID,Student Name\n441001,أحمد\n".encode("utf-8-sig")
    info = api.post(f"/courses/{cid}/roster", files={"file": ("r.csv", csv_bytes, "text/csv")}).json()
    temp = info["new_accounts"][0]["temporary_password"]
    tok = anon.post("/auth/login", json={"university_id": "441001", "password": temp}).json()
    assert tok["must_change_password"] is True
    h = {"Authorization": f"Bearer {tok['access_token']}"}
    assert anon.get("/me/courses", headers=h).status_code == 403  # must change the temporary password first
    r = anon.post("/auth/change-password", headers=h, json={"current_password": temp, "new_password": "short"})
    assert r.status_code == 400
    r = anon.post("/auth/change-password", headers=h, json={"current_password": temp, "new_password": "MyNewPass123"})
    assert r.status_code == 200
    h2 = {"Authorization": f"Bearer {r.json()['access_token']}"}
    assert anon.get("/me/courses", headers=h).status_code == 401  # old token is dead after the change
    assert anon.get("/courses", headers=h2).status_code == 403
    assert anon.post(f"/courses", headers=h2, json={"name": "x"}).status_code == 403
    assert anon.get("/admin/users", headers=h2).status_code == 403
    assert api.get("/admin/users").status_code == 403  # faculty is not admin


def test_faculty_cannot_touch_other_faculty_data(client_api):
    api, fake = client_api
    cid, eid, sids = _graded_exam(api, fake, ["441001"])
    other = TestClient(app)
    other.headers.update(api.login("F2"))
    assert other.get("/courses").json() == []
    for path in (f"/courses/{cid}", f"/courses/{cid}/exams", f"/exams/{eid}", f"/exams/{eid}/gradebook",
                 f"/exams/{eid}/submissions", f"/submissions/{sids['441001']}/review", f"/exams/{eid}/export/csv"):
        assert other.get(path).status_code == 404, path
    assert other.post(f"/submissions/{sids['441001']}/approve").status_code == 404
    assert other.post(f"/exams/{eid}/publish").status_code == 404
    assert other.delete(f"/courses/{cid}").status_code == 404
    admin = TestClient(app)
    admin.headers.update(api.login("A1"))
    assert admin.get(f"/exams/{eid}").status_code == 200  # admin oversight


def test_student_sees_only_own_approved_published_final_marks(client_api):
    api, fake = client_api
    csv_bytes = "الرقم الجامعي,اسم الطالب,البريد,الشعبة\n441001,أحمد الشمري,a@x.edu,1\n441002,فهد العنزي,b@x.edu,1\n".encode("utf-8-sig")
    cid = api.post("/courses", json={"name": "Reinforced Concrete", "code": "CE 302"}).json()["id"]
    imp = api.post(f"/courses/{cid}/roster", files={"file": ("r.csv", csv_bytes, "text/csv")}).json()
    assert imp["enrolled"] == 2 and imp["accounts_created"] == 2
    assert api.get(f"/courses/{cid}/students").json()[0]["name"] == "أحمد الشمري"
    creds = {a["university_id"]: a["temporary_password"] for a in imp["new_accounts"]}

    # importing again creates no new accounts / passwords
    again = api.post(f"/courses/{cid}/roster", files={"file": ("r.csv", csv_bytes, "text/csv")}).json()
    assert again["accounts_created"] == 0 and again["already_enrolled"] == 2 and again["new_accounts"] == []

    eid = api.post(f"/courses/{cid}/exams", data={"title": "Midterm"}).json()["id"]
    qid = api.post(f"/exams/{eid}/questions", json={"number": 1, "text": "Q", "max_mark": 10}).json()["id"]
    api.post(f"/questions/{qid}/rubric/generate?mode=default")
    api.post(f"/exams/{eid}/rubrics/approve-all")
    crit = api.get(f"/exams/{eid}/questions").json()[0]["rubric"]["criteria"]
    bd, left = [], 8.0
    for c in crit:
        m = min(c["max_mark"], left); left -= m
        bd.append({"criterion": c["criterion"], "mark": m, "max_mark": c["max_mark"]})
    sid = api.post(f"/exams/{eid}/submissions", data={"university_id": "441001"},
                   files={"files": ("p.png", b"\x89PNGx", "image/png")}).json()["id"]
    fake.messages.outputs.append({"questions": [item(1, mark=8, conf=0.95, breakdown=bd)]})
    api.post(f"/submissions/{sid}/grade")
    grade_id = api.get(f"/submissions/{sid}/review").json()["grades"][0]["id"]
    api.patch(f"/grades/{grade_id}", json={"final_mark": 6.5, "reviewer_note": "internal note"})  # AI said 8

    def student(uid):
        c = TestClient(app)
        t = c.post("/auth/login", json={"university_id": uid, "password": creds[uid]}).json()["access_token"]
        c.headers.update({"Authorization": f"Bearer {t}"})
        r = c.post("/auth/change-password", json={"current_password": creds[uid], "new_password": "StudentPass1"})
        c.headers.update({"Authorization": f"Bearer {r.json()['access_token']}"})
        return c

    s1, s2 = student("441001"), student("441002")
    assert [c["code"] for c in s1.get("/me/courses").json()] == ["CE 302"]
    assert s1.get(f"/me/courses/{cid}/grades").json() == []           # nothing until approved AND published
    api.post(f"/submissions/{sid}/approve")
    assert s1.get(f"/me/courses/{cid}/grades").json() == []           # approved but not published
    assert api.post(f"/exams/{eid}/publish").json()["visible_to_students"] == 1
    r = s1.get(f"/me/courses/{cid}/grades")
    body = r.json()
    assert body[0]["total"] == 6.5 and body[0]["questions"][0]["mark"] == 6.5
    text = r.text
    assert "ai_" not in text and "internal note" not in text and "8.0" not in text and "confidence" not in text
    assert s2.get(f"/me/courses/{cid}/grades").json() == []           # no data about others
    assert s1.get(f"/exams/{eid}").status_code == 403 and s1.get(f"/submissions/{sid}/review").status_code == 403

    # not enrolled in another instructor's course -> 404
    other = TestClient(app); other.headers.update(api.login("F2"))
    cid2 = other.post("/courses", json={"name": "Other"}).json()["id"]
    assert s1.get(f"/me/courses/{cid2}/grades").status_code == 404

    # unpublish hides again
    api.post(f"/exams/{eid}/unpublish")
    assert s1.get(f"/me/courses/{cid}/grades").json() == []

    # faculty password reset kills the old session
    sid_student = api.get(f"/courses/{cid}/students").json()[0]["student_id"]
    assert api.post(f"/courses/{cid}/students/{sid_student}/reset-password").json()["temporary_password"]
    assert s1.get("/me/courses").status_code == 401


def test_login_lockout(client_api):
    api, _ = client_api
    anon = TestClient(app)
    for _ in range(5):
        assert anon.post("/auth/login", json={"university_id": "F2", "password": "nope"}).status_code == 401
    r = anon.post("/auth/login", json={"university_id": "F2", "password": PASSWORD})
    assert r.status_code == 429


def test_audit_trail(client_api):
    api, fake = client_api
    cid, eid, sids = _graded_exam(api, fake, ["441001"], ai_mark=8)
    gid = api.get(f"/submissions/{sids['441001']}/review").json()["grades"][0]["id"]
    api.patch(f"/grades/{gid}", json={"final_mark": 7, "reviewer_note": "alt. method"})
    api.post(f"/submissions/{sids['441001']}/approve")
    api.post(f"/exams/{eid}/publish")
    admin = TestClient(app); admin.headers.update(api.login("A1"))
    log = admin.get("/admin/audit?limit=100").json()
    actions = [a["action"] for a in log]
    assert {"login", "course_created", "grade_set", "submission_approved", "grades_published"} <= set(actions)
    gs = next(a for a in log if a["action"] == "grade_set")
    assert gs["actor"] == "F1" and gs["details"]["ai_mark"] == 8 and gs["details"]["after"] == 7 and gs["details"]["before"] is None
    assert "password" not in json.dumps(log).lower().replace("password_changed", "").replace("password_reset", "")
    assert api.get("/admin/audit").status_code == 403


def test_admin_creates_users(client_api):
    api, _ = client_api
    admin = TestClient(app); admin.headers.update(api.login("A1"))
    r = admin.post("/admin/users", json={"university_id": "F9", "name": "Dr. New", "role": "faculty"})
    assert r.status_code == 201 and len(r.json()["temporary_password"]) >= 10
    assert admin.post("/admin/users", json={"university_id": "F9", "name": "Dup", "role": "faculty"}).status_code == 409
    row = admin.get("/admin/users").json()
    assert all("password" not in k for u in row for k in u)


# ------------------------------------------------------------------ review requests (appeals) + notifications


def _breakdown(crit, mark):
    out, left = [], mark
    for c in crit:
        m = min(c["max_mark"], left)
        left -= m
        out.append({"criterion": c["criterion"], "mark": m, "max_mark": c["max_mark"]})
    return out


def _student(uid, temp):
    c = TestClient(app)
    t = c.post("/auth/login", json={"university_id": uid, "password": temp}).json()["access_token"]
    r = c.post("/auth/change-password", headers={"Authorization": f"Bearer {t}"},
               json={"current_password": temp, "new_password": "StudentPass1"})
    c.headers.update({"Authorization": f"Bearer {r.json()['access_token']}"})
    return c


def _published_world(api, fake, publish=True, appeal_days=7):
    """441001 (final 8/10) and 441002 (9/10) approved; 441003 graded but NOT approved. Exam published."""
    csv_bytes = "University ID,Student Name\n441001,أحمد الشمري\n441002,فهد العنزي\n441003,سارة\n".encode("utf-8-sig")
    cid = api.post("/courses", json={"name": "Reinforced Concrete", "code": "CE 302"}).json()["id"]
    imp = api.post(f"/courses/{cid}/roster", files={"file": ("r.csv", csv_bytes, "text/csv")}).json()
    creds = {a["university_id"]: a["temporary_password"] for a in imp["new_accounts"]}
    eid = api.post(f"/courses/{cid}/exams", data={"title": "Midterm"}).json()["id"]
    qid = api.post(f"/exams/{eid}/questions", json={"number": 1, "text": "Q1", "max_mark": 10}).json()["id"]
    api.post(f"/questions/{qid}/rubric/generate?mode=default")
    api.post(f"/exams/{eid}/rubrics/approve-all")
    crit = api.get(f"/exams/{eid}/questions").json()[0]["rubric"]["criteria"]
    sids = {}
    for uid, mark in (("441001", 8.0), ("441002", 9.0), ("441003", 7.0)):
        sid = api.post(f"/exams/{eid}/submissions", data={"university_id": uid},
                       files={"files": ("p.png", b"\x89PNG-" + uid.encode(), "image/png")}).json()["id"]
        fake.messages.outputs.append({"questions": [item(1, mark=mark, conf=0.95, breakdown=_breakdown(crit, mark))]})
        assert api.post(f"/submissions/{sid}/grade").status_code == 200
        sids[uid] = sid
    for uid in ("441001", "441002"):
        assert api.post(f"/submissions/{sids[uid]}/approve").status_code == 200
    if publish:
        assert api.post(f"/exams/{eid}/publish?appeal_days={appeal_days}").status_code == 200
    return {"cid": cid, "eid": eid, "sids": sids, "creds": creds, "crit": crit}


def test_appeal_full_flow(client_api):
    api, fake = client_api
    w = _published_world(api, fake)
    s1, s2 = _student("441001", w["creds"]["441001"]), _student("441002", w["creds"]["441002"])
    g = s1.get(f"/me/courses/{w['cid']}/grades").json()[0]
    assert g["review"]["can_request"] is True and g["total"] == 8

    # --- student files a request
    r = s1.post(f"/me/exams/{w['eid']}/review-requests",
                json={"question_numbers": [1], "message": "My method is correct; the arithmetic slip should cost less."})
    assert r.status_code == 201, r.text
    rid = r.json()["id"]
    assert "ai_" not in r.text and r.json()["status"] == "pending"
    assert s1.post(f"/me/exams/{w['eid']}/review-requests",
                   json={"question_numbers": [1], "message": "second attempt should fail"}).status_code == 409
    assert s1.get(f"/me/courses/{w['cid']}/grades").json()[0]["review"]["can_request"] is False

    # --- privacy: other student / other instructor see nothing
    assert s2.get("/me/review-requests").json() == []
    assert s2.get(f"/me/review-requests/{rid}/paper").status_code == 404
    other = TestClient(app); other.headers.update(api.login("F2"))
    assert other.get("/review-requests").json() == [] and other.get(f"/review-requests/{rid}").status_code == 404
    assert other.post(f"/review-requests/{rid}/decide", json={"items": [], "response_message": "x"}).status_code in (404, 422)
    assert s1.get("/review-requests").status_code == 403  # students have no faculty endpoints

    # --- instructor is notified and sees the case
    n = api.get("/notifications?unread=true").json()
    assert n["unread"] == 1 and n["items"][0]["kind"] == "review_requested"
    listing = api.get("/review-requests?status=pending").json()
    assert [x["id"] for x in listing] == [rid] and listing[0]["questions"] == [1]
    d = api.get(f"/review-requests/{rid}").json()
    assert d["message"].startswith("My method") and d["items"][0]["original_mark"] == 8 and d["items"][0]["second_opinion"] is None
    assert api.post(f"/notifications/{n['items'][0]['id']}/read").status_code == 200
    assert api.get("/notifications").json()["unread"] == 0

    # --- AI second opinion: independent, advisory, prompt-injection aware, changes nothing
    fake.messages.outputs.append({"questions": [item(1, mark=9, conf=0.88, breakdown=_breakdown(w["crit"], 9), reason="Alt. method earns method mark")]})
    a = api.post(f"/review-requests/{rid}/ai-assist").json()
    so = a["items"][0]["second_opinion"]
    assert so["mark"] == 9 and so["difference"] == 1 and so["flags"] == []
    call = fake.messages.calls[-1]
    texts = " ".join(b.get("text", "") for b in call["messages"][0]["content"] if b["type"] == "text")
    assert "SECOND OPINION" in call["system"] and "UNTRUSTED" in call["system"]
    assert "STUDENT'S OBJECTION (untrusted" in texts and "My method is correct" in texts
    assert s1.get(f"/me/courses/{w['cid']}/grades").json()[0]["total"] == 8  # still 8: AI never changes marks

    # --- decision validation
    qid = d["items"][0]["question_id"]
    bad = lambda **kw: api.post(f"/review-requests/{rid}/decide", json={"response_message": "ok", **kw})
    assert bad(items=[{"question_id": qid + 999, "decision": "keep"}]).status_code == 400
    assert bad(items=[{"question_id": qid, "decision": "change", "new_mark": 11}]).status_code == 400
    assert bad(items=[{"question_id": qid, "decision": "change", "new_mark": 8}]).status_code == 400
    assert bad(items=[{"question_id": qid, "decision": "change"}]).status_code == 400
    assert api.post(f"/review-requests/{rid}/decide", json={"items": [{"question_id": qid, "decision": "keep"}], "response_message": ""}).status_code == 422
    assert s1.get(f"/me/review-requests/{rid}/paper").status_code == 403  # no access before a decision

    # --- decide: change 8 -> 9, release the paper
    ok = bad(items=[{"question_id": qid, "decision": "change", "new_mark": 9, "note": "Method credit restored"}],
             response_message="Thanks - you were right about the method mark.", grant_paper_access=True)
    assert ok.status_code == 200 and ok.json()["status"] == "decided"
    assert bad(items=[{"question_id": qid, "decision": "keep"}]).status_code == 409  # cannot decide twice

    mine = s1.get("/me/review-requests").json()[0]
    assert mine["status"] == "decided" and mine["items"][0]["new_mark"] == 9 and mine["items"][0]["note"] == "Method credit restored"
    assert mine["response_message"].startswith("Thanks") and "ai_" not in json.dumps(mine)
    assert s1.get(f"/me/courses/{w['cid']}/grades").json()[0]["total"] == 9
    sn = s1.get("/notifications").json()
    assert sn["unread"] == 1 and sn["items"][0]["kind"] == "review_decided"

    paper = s1.get(f"/me/review-requests/{rid}/paper").json()
    assert paper["total"] == 9 and paper["questions"][0]["mark"] == 9 and "ai_" not in json.dumps(paper)
    f = s1.get(f"/me/review-requests/{rid}/file/0")
    assert f.status_code == 200 and f.content.startswith(b"\x89PNG-441001")
    assert s2.get(f"/me/review-requests/{rid}/file/0").status_code == 404
    api.post(f"/review-requests/{rid}/paper-access?grant=false")
    assert s1.get(f"/me/review-requests/{rid}/paper").status_code == 403

    # --- audit: before/after and the AI opinion are recorded
    admin = TestClient(app); admin.headers.update(api.login("A1"))
    ev = next(a for a in admin.get("/admin/audit?action=review_decided").json())
    assert ev["details"]["items"][0] == {"question": 1, "before": 8.0, "after": 9.0, "ai_second_opinion": 9.0}
    assert {"review_requested", "review_ai_assist", "own_paper_viewed"} <= {a["action"] for a in admin.get("/admin/audit?limit=200").json()}
    assert api.get(f"/exams/{w['eid']}/gradebook").json()["rows"][0]["total"] == 9  # gradebook follows the decision


def test_appeal_rules(client_api):
    api, fake = client_api
    w = _published_world(api, fake, publish=False)
    s1, s3 = _student("441001", w["creds"]["441001"]), _student("441003", w["creds"]["441003"])
    body = {"question_numbers": [1], "message": "Please re-check question one, thank you."}
    assert s1.post(f"/me/exams/{w['eid']}/review-requests", json=body).status_code == 404   # not published -> invisible
    assert api.post(f"/exams/{w['eid']}/publish?appeal_days=100").status_code == 400
    api.post(f"/exams/{w['eid']}/publish?appeal_days=0")
    assert s1.post(f"/me/exams/{w['eid']}/review-requests", json=body).status_code == 409   # window closed
    assert s1.get(f"/me/courses/{w['cid']}/grades").json()[0]["review"]["can_request"] is False
    api.post(f"/exams/{w['eid']}/publish?appeal_days=3")
    assert s3.post(f"/me/exams/{w['eid']}/review-requests", json=body).status_code == 409   # paper not approved -> no result yet
    assert s1.post(f"/me/exams/{w['eid']}/review-requests", json={**body, "message": "short"}).status_code == 422
    assert s1.post(f"/me/exams/{w['eid']}/review-requests", json={**body, "question_numbers": [7]}).status_code == 400
    assert s1.post(f"/me/exams/{w['eid']}/review-requests", json={**body, "question_numbers": []}).status_code == 422
    # window expiry
    db = api.S()
    from datetime import datetime, timedelta, timezone
    db.get(Exam, w["eid"]).appeals_deadline = datetime.now(timezone.utc) - timedelta(minutes=1)
    db.commit(); db.close()
    assert s1.post(f"/me/exams/{w['eid']}/review-requests", json=body).status_code == 409
    # faculty endpoints are closed to students, student endpoints to faculty
    assert s1.post("/review-requests/1/ai-assist").status_code == 403
    assert api.get("/me/review-requests").status_code == 403


def test_appeal_reject(client_api):
    api, fake = client_api
    w = _published_world(api, fake)
    s1 = _student("441001", w["creds"]["441001"])
    rid = s1.post(f"/me/exams/{w['eid']}/review-requests",
                  json={"question_numbers": [1], "message": "Please look again at my answer."}).json()["id"]
    assert api.post(f"/review-requests/{rid}/reject", json={"response_message": ""}).status_code == 422
    r = api.post(f"/review-requests/{rid}/reject", json={"response_message": "The rubric was applied correctly."}).json()
    assert r["status"] == "rejected"
    mine = s1.get("/me/review-requests").json()[0]
    assert mine["status"] == "rejected" and mine["response_message"].startswith("The rubric") and mine["paper_access"] is False
    assert s1.get(f"/me/review-requests/{rid}/paper").status_code == 403
    assert api.post(f"/review-requests/{rid}/reject", json={"response_message": "again"}).status_code == 409
    assert api.post(f"/review-requests/{rid}/paper-access?grant=true").status_code == 409   # only after a decision
    assert s1.post(f"/me/exams/{w['eid']}/review-requests",
                   json={"question_numbers": [1], "message": "Trying again with a new request."}).status_code == 409
    assert s1.get(f"/me/courses/{w['cid']}/grades").json()[0]["total"] == 8  # marks untouched
    assert s1.get("/notifications").json()["items"][0]["kind"] == "review_rejected"


# ------------------------------------------------------------------ bulk grade updates via Excel


def _graded_world_for_bulk(api, fake):
    """Two students, one question, 8/10 each; one approved, one not."""
    cid = api.post("/courses", json={"name": "Reinforced Concrete", "code": "CE 302"}).json()["id"]
    eid = api.post(f"/courses/{cid}/exams", data={"title": "Midterm"}).json()["id"]
    qid = api.post(f"/exams/{eid}/questions", json={"number": 1, "text": "Q1", "max_mark": 10}).json()["id"]
    api.post(f"/questions/{qid}/rubric/generate?mode=default")
    api.post(f"/exams/{eid}/rubrics/approve-all")
    crit = api.get(f"/exams/{eid}/questions").json()[0]["rubric"]["criteria"]
    sids = {}
    for uid in ("441001", "441002"):
        sid = api.post(f"/exams/{eid}/submissions", data={"university_id": uid},
                       files={"files": ("p.png", b"\x89PNGx", "image/png")}).json()["id"]
        fake.messages.outputs.append({"questions": [item(1, mark=8, conf=0.95, breakdown=_breakdown(crit, 8))]})
        assert api.post(f"/submissions/{sid}/grade").status_code == 200
        sids[uid] = sid
    api.post(f"/submissions/{sids['441001']}/approve")
    return {"cid": cid, "eid": eid, "qid": qid, "sids": sids}


def test_editable_export_is_reuploadable_and_sheet_protected(client_api):
    import openpyxl
    api, fake = client_api
    w = _graded_world_for_bulk(api, fake)
    r = api.get(f"/exams/{w['eid']}/gradebook/editable-xlsx")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert ws.protection.sheet is True
    assert ws["A2"].value == "University ID" and ws["C2"].value.startswith("Q1")
    ids = {ws.cell(row, 1).value for row in (3, 4)}
    assert ids == {"441001", "441002"}
    assert ws.cell(3, 1).protection.locked is True   # id column stays locked
    assert ws.cell(3, 3).protection.locked is False  # mark cell is editable


def test_bulk_import_preview_then_commit(client_api):
    import openpyxl
    api, fake = client_api
    w = _graded_world_for_bulk(api, fake)
    xlsx = openpyxl.load_workbook(io.BytesIO(api.get(f"/exams/{w['eid']}/gradebook/editable-xlsx").content))
    ws = xlsx.active
    for row in (3, 4):
        if ws.cell(row, 1).value == "441001":
            ws.cell(row, 3).value = 9.5  # was approved at 8 -> should require reopen
        else:
            ws.cell(row, 3).value = 6.0  # not approved, plain edit
    buf = io.BytesIO()
    xlsx.save(buf)
    upload = ("edited.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    preview = api.post(f"/exams/{w['eid']}/gradebook/import", files={"file": upload}).json()
    assert preview["committed"] is False and len(preview["changes"]) == 2
    by_uid = {c["university_id"]: c for c in preview["changes"]}
    assert by_uid["441001"]["current_mark"] == 8 and by_uid["441001"]["new_mark"] == 9.5 and by_uid["441001"]["currently_approved"] is True
    assert by_uid["441002"]["new_mark"] == 6.0 and by_uid["441002"]["currently_approved"] is False
    # preview must not have written anything
    assert api.get(f"/submissions/{w['sids']['441001']}/review").json()["grades"][0]["final_mark"] == 8

    commit = api.post(f"/exams/{w['eid']}/gradebook/import", params={"commit": True}, files={"file": upload}).json()
    assert commit["committed"] is True and commit["applied"] == 2 and commit["reopened_submissions"] == 1

    rev1 = api.get(f"/submissions/{w['sids']['441001']}/review").json()
    assert rev1["status"] == "graded"  # reopened, needs re-approval
    assert rev1["grades"][0]["final_mark"] == 9.5 and rev1["grades"][0]["decision"] == "edited"
    assert "Bulk Excel import" in (rev1["grades"][0]["reviewer_note"] or "")
    rev2 = api.get(f"/submissions/{w['sids']['441002']}/review").json()
    assert rev2["grades"][0]["final_mark"] == 6.0

    # re-uploading the same file again is a no-op (already applied)
    again = api.post(f"/exams/{w['eid']}/gradebook/import", files={"file": upload}).json()
    assert again["changes"] == []

    admin = TestClient(app); admin.headers.update(api.login("A1"))
    log = admin.get("/admin/audit?limit=200").json()
    actions = [a["action"] for a in log]
    assert {"submission_reopened", "grade_set", "bulk_excel_import_committed"} <= set(actions)
    gs = [a for a in log if a["action"] == "grade_set" and a["details"].get("source") == "bulk_excel_import"]
    assert any(g["details"]["before"] == 8 and g["details"]["after"] == 9.5 for g in gs)


def test_bulk_import_rejects_out_of_range_and_unknown_student(client_api):
    import openpyxl
    api, fake = client_api
    w = _graded_world_for_bulk(api, fake)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["University ID", "Student", "Q1 (max 10)"])
    ws.append(["441001", "x", 99])       # out of range
    ws.append(["999999", "ghost", 5])    # not in this exam
    buf = io.BytesIO(); wb.save(buf)
    upload = ("bad.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    r = api.post(f"/exams/{w['eid']}/gradebook/import", files={"file": upload}).json()
    assert r["changes"] == []
    assert any("outside" in wtext for wtext in r["warnings"])
    assert r["unmatched"] and r["unmatched"][0]["university_id"] == "999999"


def test_bulk_import_ownership_and_bad_file(client_api):
    api, fake = client_api
    w = _graded_world_for_bulk(api, fake)
    other = TestClient(app); other.headers.update(api.login("F2"))
    upload = ("x.xlsx", b"not an excel file", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert other.get(f"/exams/{w['eid']}/gradebook/editable-xlsx").status_code == 404
    assert other.post(f"/exams/{w['eid']}/gradebook/import", files={"file": upload}).status_code == 404
    assert api.post(f"/exams/{w['eid']}/gradebook/import", files={"file": upload}).status_code == 400


# ------------------------------------------------------------------ assignments


def _setup_assignment(api, fake, due_at=None, allow_late=False, max_mark=10):
    csv_bytes = "University ID,Student Name\n441001,أحمد الشمري\n441002,فهد العنزي\n".encode("utf-8-sig")
    cid = api.post("/courses", json={"name": "Reinforced Concrete", "code": "CE 302"}).json()["id"]
    imp = api.post(f"/courses/{cid}/roster", files={"file": ("r.csv", csv_bytes, "text/csv")}).json()
    creds = {a["university_id"]: a["temporary_password"] for a in imp["new_accounts"]}
    body = {"title": "Homework 1", "description": "Solve problems 1-3", "max_mark": max_mark, "allow_late": allow_late}
    if due_at:
        body["due_at"] = due_at
    aid = api.post(f"/courses/{cid}/assignments", json=body).json()["id"]
    return {"cid": cid, "aid": aid, "creds": creds}


def test_assignment_grading_requires_approved_rubric(client_api):
    api, fake = client_api
    w = _setup_assignment(api, fake)
    s1 = _student("441001", w["creds"]["441001"])
    r = s1.post(f"/me/assignments/{w['aid']}/submit", files={"files": ("hw.png", b"\x89PNGx", "image/png")})
    assert r.status_code == 201 and r.json()["status"] == "uploaded"
    assert api.post(f"/assignment-submissions/{r.json()['id']}/grade").status_code == 409  # no rubric yet

    api.post(f"/assignments/{w['aid']}/rubric/generate?mode=default")
    rub = api.get(f"/assignments/{w['aid']}").json()["rubric"]
    ok = api.put(f"/assignments/{w['aid']}/rubric", json=rub)
    assert ok.status_code == 200 and ok.json()["rubric_approved"] is True

    fake.messages.outputs.append({"questions": [item(1, mark=8, conf=0.92, breakdown=_breakdown(rub["criteria"], 8))]})
    g = api.post(f"/assignment-submissions/{r.json()['id']}/grade")
    assert g.status_code == 200 and g.json()["ai_mark"] == 8 and g.json()["status"] == "graded"


def test_assignment_full_flow_submit_grade_approve_release(client_api):
    api, fake = client_api
    w = _setup_assignment(api, fake)
    api.post(f"/assignments/{w['aid']}/rubric/generate?mode=default")
    rub = api.get(f"/assignments/{w['aid']}").json()["rubric"]
    api.put(f"/assignments/{w['aid']}/rubric", json=rub)

    s1 = _student("441001", w["creds"]["441001"])
    sid = s1.post(f"/me/assignments/{w['aid']}/submit", files={"files": ("hw.png", b"\x89PNGx", "image/png")}).json()["id"]
    # resubmitting before grading overwrites
    assert s1.post(f"/me/assignments/{w['aid']}/submit", files={"files": ("hw2.png", b"\x89PNGy", "image/png")}).status_code == 201

    fake.messages.outputs.append({"questions": [item(1, mark=6, conf=0.4, breakdown=_breakdown(rub["criteria"], 6))]})
    assert api.post(f"/assignment-submissions/{sid}/grade").status_code == 200
    # cannot resubmit once graded
    assert s1.post(f"/me/assignments/{w['aid']}/submit", files={"files": ("x.png", b"\x89PNGz", "image/png")}).status_code == 409

    review = api.get(f"/assignment-submissions/{sid}/review").json()
    assert review["needs_review"] is True  # low confidence
    assert api.post(f"/assignment-submissions/{sid}/approve").status_code == 409  # flagged, must decide first
    patched = api.patch(f"/assignment-submissions/{sid}", json={"final_mark": 7, "reviewer_note": "partial credit"})
    assert patched.status_code == 200 and patched.json()["decision"] == "edited"
    assert api.post(f"/assignment-submissions/{sid}/approve").json()["status"] == "approved"

    # not visible to student until released
    before = s1.get(f"/me/courses/{w['cid']}/assignments").json()[0]
    assert before["status"] == "approved" and before["mark"] is None
    api.post(f"/assignments/{w['aid']}/release")
    after = s1.get(f"/me/courses/{w['cid']}/assignments").json()[0]
    assert after["mark"] == 7

    admin = TestClient(app); admin.headers.update(api.login("A1"))
    actions = {a["action"] for a in admin.get("/admin/audit?limit=200").json()}
    assert {"assignment_submitted", "assignment_grade_set", "assignment_submission_approved",
            "assignment_results_released"} <= actions


def test_assignment_due_date_and_late_rules(client_api):
    from datetime import datetime, timedelta, timezone
    api, fake = client_api
    past = (datetime.now(timezone.utc) - timedelta(days=1)).isoformat()
    w = _setup_assignment(api, fake, due_at=past, allow_late=False)
    s1 = _student("441001", w["creds"]["441001"])
    r = s1.post(f"/me/assignments/{w['aid']}/submit", files={"files": ("hw.png", b"\x89PNGx", "image/png")})
    assert r.status_code == 409

    w2 = _setup_assignment(api, fake, due_at=past, allow_late=True)
    s2 = _student("441002", w["creds"]["441002"])  # account already exists from w's roster import
    r2 = s2.post(f"/me/assignments/{w2['aid']}/submit", files={"files": ("hw.png", b"\x89PNGx", "image/png")})
    assert r2.status_code == 201 and r2.json()["late"] is True


def test_assignment_ownership_and_privacy(client_api):
    api, fake = client_api
    w = _setup_assignment(api, fake)
    other = TestClient(app); other.headers.update(api.login("F2"))
    assert other.get(f"/assignments/{w['aid']}").status_code == 404
    assert other.get(f"/courses/{w['cid']}/assignments").status_code == 404
    s1 = _student("441001", w["creds"]["441001"])
    assert s1.post(f"/courses/{w['cid']}/assignments", json={"title": "x", "max_mark": 5}).status_code == 403
    assert api.get("/me/courses/1/assignments").status_code == 403  # faculty has no student endpoints


# ------------------------------------------------------------------ self-service registration & password reset


def test_faculty_self_registration(client_api):
    api, fake = client_api
    anon = TestClient(app)
    r = anon.post("/auth/register", json={"university_id": "F5", "name": "Dr. New", "email": "new@nbu.edu.sa",
                                          "password": "MyOwnPass1"})
    assert r.status_code == 201
    body = r.json()
    assert body["role"] == "faculty" and body["must_change_password"] is False
    h = {"Authorization": f"Bearer {body['access_token']}"}
    assert anon.get("/courses", headers=h).json() == []  # sees only own courses (none yet)
    created = anon.post("/courses", headers=h, json={"name": "My Own Course"})
    assert created.status_code == 201

    # duplicate id / email rejected
    dup1 = anon.post("/auth/register", json={"university_id": "F5", "name": "x", "email": "other@nbu.edu.sa", "password": "AnotherPass1"})
    assert dup1.status_code == 409
    dup2 = anon.post("/auth/register", json={"university_id": "F6", "name": "x", "email": "new@nbu.edu.sa", "password": "AnotherPass1"})
    assert dup2.status_code == 409

    # too-short password / bad email rejected
    assert anon.post("/auth/register", json={"university_id": "F7", "name": "x", "email": "a@b.com", "password": "short"}).status_code == 422
    assert anon.post("/auth/register", json={"university_id": "F8", "name": "x", "email": "not-an-email", "password": "GoodPass1"}).status_code == 422

    # audited
    admin = TestClient(app); admin.headers.update(api.login("A1"))
    actions = {a["action"] for a in admin.get("/admin/audit?limit=200").json()}
    assert "faculty_self_registered" in actions


def test_forgot_and_reset_password_flow(client_api, monkeypatch):
    api, fake = client_api
    anon = TestClient(app)
    anon.post("/auth/register", json={"university_id": "F5", "name": "Dr. New", "email": "reset@nbu.edu.sa", "password": "OldPassw0rd"})

    sent = {}
    import app.routers.auth as auth_mod
    monkeypatch.setattr(auth_mod, "send_email", lambda to, subject, body: sent.update(to=to, subject=subject, body=body))

    r1 = anon.post("/auth/forgot-password", json={"university_id": "F5"})
    assert r1.status_code == 200 and "reset link has been sent" in r1.json()["message"]
    assert sent["to"] == "reset@nbu.edu.sa" and "reset_token=" in sent["body"]
    token = sent["body"].split("reset_token=")[1].split()[0].strip()

    # unknown account: identical generic response (no user enumeration), no email attempt recorded
    sent.clear()
    r2 = anon.post("/auth/forgot-password", json={"university_id": "NOPE"})
    assert r2.json() == r1.json() and sent == {}

    # wrong/garbage token rejected
    assert anon.post("/auth/reset-password", json={"token": "garbage-token-value", "new_password": "NewPassw0rd"}).status_code == 400

    # correct token resets the password and logs the user in
    ok = anon.post("/auth/reset-password", json={"token": token, "new_password": "NewPassw0rd"})
    assert ok.status_code == 200
    assert anon.post("/auth/login", json={"university_id": "F5", "password": "OldPassw0rd"}).status_code == 401
    assert anon.post("/auth/login", json={"university_id": "F5", "password": "NewPassw0rd"}).status_code == 200

    # token is single-use
    reuse = anon.post("/auth/reset-password", json={"token": token, "new_password": "AnotherPass2"})
    assert reuse.status_code == 400

    admin = TestClient(app); admin.headers.update(api.login("A1"))
    actions = {a["action"] for a in admin.get("/admin/audit?limit=200").json()}
    assert {"password_reset_requested", "password_reset_completed"} <= actions


def test_forgot_password_account_with_no_email_is_silent(client_api):
    api, fake = client_api
    anon = TestClient(app)
    # F1 was seeded without an email in the fixture
    r = anon.post("/auth/forgot-password", json={"university_id": "F1"})
    assert r.status_code == 200 and "reset link has been sent" in r.json()["message"]
